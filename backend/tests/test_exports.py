"""T6.1-T6.4 — TXT/DOCX/PDF rendered from the canonical transcript."""

from datetime import UTC, datetime

import pytest
from docx import Document
from reportlab.pdfbase import pdfmetrics

from app.core.errors import AppError, ErrorCode
from app.schemas.transcript import (
    EngineInfo,
    Transcript,
    TranscriptSegment,
    TranscriptSource,
    TranscriptSpeaker,
    TranscriptWord,
)
from app.services import exports

CYRILLIC = "Добрый день. Давайте обсудим результаты."


def make_transcript(segments=None, duration=3522.41) -> Transcript:
    return Transcript(
        job_id="11111111-1111-1111-1111-111111111111",
        source=TranscriptSource(filename="meeting.m4a", duration_seconds=duration),
        asr=EngineInfo(engine="gigaam", model="v3_e2e_rnnt"),
        diarization=EngineInfo(engine="pyannote", model="pyannote/speaker-diarization-community-1"),
        speakers=[
            TranscriptSpeaker(id="SPEAKER_00", display_name="Спикер 1"),
            TranscriptSpeaker(id="SPEAKER_01", display_name="Анна"),
        ],
        words=[TranscriptWord(start=12.42, end=12.83, text="Добрый", speaker_id="SPEAKER_00")],
        segments=[
            TranscriptSegment(
                id="seg-001", start=12.42, end=18.73, speaker_id="SPEAKER_00", text=CYRILLIC
            ),
            TranscriptSegment(
                id="seg-002",
                start=18.73,
                end=27.0,
                speaker_id="SPEAKER_01",
                text="Я подготовил данные.",
            ),
        ]
        if segments is None
        else segments,
        created_at=datetime.now(UTC),
    )


@pytest.fixture
def transcript() -> Transcript:
    return make_transcript()


def font_or_skip():
    try:
        return exports.resolve_pdf_font(None)
    except AppError:
        pytest.skip("no Cyrillic-capable TTF on this machine")


# --- timestamps ----------------------------------------------------------------------


def test_timestamp_formatting() -> None:
    assert exports.format_timestamp(12.42) == "00:00:12"
    assert exports.format_timestamp(3522.41) == "00:58:42"
    assert exports.format_timestamp(3661) == "01:01:01"
    assert exports.format_timestamp(None) == "00:00:00"


# --- T6.1 TXT ------------------------------------------------------------------------


def test_txt_matches_the_spec_layout(transcript) -> None:
    text = exports.render_txt(transcript)

    assert text.startswith("[00:00:12 – 00:00:18] Спикер 1\n")
    assert CYRILLIC in text
    assert "\n\n[00:00:18 – 00:00:27] Анна\n" in text


def test_txt_uses_the_renamed_speaker(transcript) -> None:
    assert "Анна" in exports.render_txt(transcript)
    assert "Спикер 2" not in exports.render_txt(transcript)


def test_txt_labels_unattributed_speech(tmp_path) -> None:
    transcript = make_transcript(
        segments=[TranscriptSegment(id="seg-001", start=0.0, end=1.0, speaker_id=None, text="Шум")]
    )

    assert exports.UNKNOWN_SPEAKER in exports.render_txt(transcript)


def test_txt_file_is_utf8(transcript, tmp_path) -> None:
    path = exports.write_txt(transcript, tmp_path / "transcript.txt")

    assert CYRILLIC in path.read_text(encoding="utf-8")


def test_empty_transcript_produces_an_empty_txt() -> None:
    assert exports.render_txt(make_transcript(segments=[])) == ""


# --- T6.2 DOCX -----------------------------------------------------------------------


def test_docx_contains_title_filename_duration_and_blocks(transcript, tmp_path) -> None:
    path = exports.write_docx(transcript, tmp_path / "transcript.docx")

    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    assert exports.DOCUMENT_TITLE in paragraphs
    assert "Файл: meeting.m4a" in paragraphs
    assert "Длительность: 00:58:42" in paragraphs
    assert "[00:00:12 – 00:00:18] Спикер 1" in paragraphs
    assert CYRILLIC in paragraphs


def test_docx_timestamp_heading_is_bold(transcript, tmp_path) -> None:
    path = exports.write_docx(transcript, tmp_path / "transcript.docx")

    headings = [p for p in Document(str(path)).paragraphs if p.text.startswith("[00:00:12")]
    assert headings and headings[0].runs[0].bold is True


def test_docx_without_duration_still_renders(tmp_path) -> None:
    transcript = make_transcript(duration=None)

    path = exports.write_docx(transcript, tmp_path / "transcript.docx")

    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    assert not any(text.startswith("Длительность") for text in paragraphs)


# --- T6.3 PDF ------------------------------------------------------------------------


def test_pdf_is_generated_with_an_embedded_cyrillic_font(transcript, tmp_path) -> None:
    font = font_or_skip()

    path = exports.write_pdf(transcript, tmp_path / "transcript.pdf", font)

    payload = path.read_bytes()
    assert payload.startswith(b"%PDF")
    assert payload.count(b"/FontFile2") >= 1, "the TTF must be embedded, not referenced"

    # T6.3 — the built-in PDF fonts are Latin-only, so the chosen face must really carry
    # Cyrillic glyphs, otherwise the text renders as empty boxes.
    face = pdfmetrics.getFont(exports.PDF_FONT_NAME).face
    assert all(face.charToGlyph.get(ord(letter)) for letter in "ДобрыйЁж")
    assert font.is_file()


def test_pdf_font_can_be_configured(transcript, tmp_path) -> None:
    font = font_or_skip()

    assert exports.resolve_pdf_font(font) == font


def test_missing_configured_font_fails_loudly(transcript, tmp_path) -> None:
    with pytest.raises(AppError) as exc:
        exports.resolve_pdf_font(tmp_path / "nope.ttf")

    assert exc.value.code is ErrorCode.EXPORT_FAILED


def test_no_font_anywhere_fails_loudly(monkeypatch) -> None:
    monkeypatch.setattr(exports, "FONT_CANDIDATES", ())

    with pytest.raises(AppError) as exc:
        exports.resolve_pdf_font(None)

    assert exc.value.code is ErrorCode.EXPORT_FAILED
    assert "PDF_FONT_PATH" in exc.value.message


def test_pdf_escapes_markup_in_the_transcript(tmp_path) -> None:
    """Reportlab parses paragraph markup — raw `<` from speech must not break the build."""
    font = font_or_skip()
    transcript = make_transcript(
        segments=[
            TranscriptSegment(
                id="seg-001", start=0.0, end=1.0, speaker_id=None, text="цена < 5 & больше"
            )
        ]
    )

    path = exports.write_pdf(transcript, tmp_path / "transcript.pdf", font)

    assert path.stat().st_size > 0


# --- T6.4 regeneration ---------------------------------------------------------------


def test_generate_exports_writes_all_three_views(transcript, tmp_path) -> None:
    font_or_skip()

    written = exports.generate_exports(transcript, tmp_path)

    assert [path.name for path in written] == [
        "transcript.txt",
        "transcript.docx",
        "transcript.pdf",
    ]
    assert all(path.stat().st_size > 0 for path in written)


def test_regeneration_reads_only_the_canonical_json(transcript, tmp_path) -> None:
    font_or_skip()
    source = tmp_path / "transcript.json"
    source.write_text(transcript.model_dump_json(), encoding="utf-8")

    exports.regenerate_exports(source, tmp_path)
    first = (tmp_path / "transcript.txt").read_text(encoding="utf-8")

    # Rename a speaker in the canonical JSON only — no ASR, no diarization.
    renamed = transcript.model_copy(deep=True)
    renamed.speakers[0].display_name = "Сергей"
    source.write_text(renamed.model_dump_json(), encoding="utf-8")
    exports.regenerate_exports(source, tmp_path)
    second = (tmp_path / "transcript.txt").read_text(encoding="utf-8")

    assert "Спикер 1" in first
    assert "Сергей" in second
    assert "Спикер 1" not in second


def test_regeneration_without_the_canonical_json_fails(tmp_path) -> None:
    with pytest.raises(AppError) as exc:
        exports.regenerate_exports(tmp_path / "missing.json", tmp_path)

    assert exc.value.code is ErrorCode.EXPORT_FAILED


def test_regeneration_with_a_corrupted_json_fails(tmp_path) -> None:
    source = tmp_path / "transcript.json"
    source.write_text("{not json", encoding="utf-8")

    with pytest.raises(AppError) as exc:
        exports.regenerate_exports(source, tmp_path)

    assert exc.value.code is ErrorCode.EXPORT_FAILED


def test_export_failure_has_a_stable_code(transcript, tmp_path, monkeypatch) -> None:
    def boom(*args, **kwargs):
        raise RuntimeError("disk on fire")

    monkeypatch.setattr(exports, "write_txt", boom)

    with pytest.raises(AppError) as exc:
        exports.generate_exports(transcript, tmp_path)

    assert exc.value.code is ErrorCode.EXPORT_FAILED
