"""TXT / DOCX / PDF views of the canonical transcript (TASKS.md T6.1-T6.4).

Every export is generated from `result/transcript.json` and nothing else — DOCX or PDF
must never be the only copy of a result (AGENTS.md rule 5). Renaming a speaker later
re-runs exactly this code; it never re-runs diarization or ASR.

Lives in `app/services/` rather than `worker/` because the API image regenerates exports
after a rename and cannot import anything from `worker/`.
"""

import logging
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.core.errors import AppError, ErrorCode
from app.schemas.transcript import Transcript

logger = logging.getLogger(__name__)

EXPORT_FORMATS: tuple[str, ...] = ("txt", "docx", "pdf")

DOCUMENT_TITLE = "Транскрипт"
# Shown instead of a name when a word could not be attributed (SPEC.md §3.4).
UNKNOWN_SPEAKER = "Спикер не определён"
# En dash with spaces, exactly as in the SPEC.md §6 example.
TIME_RANGE_SEPARATOR = " – "

# Where a Cyrillic-capable TTF is looked for when PDF_FONT_PATH is not set. The container
# gets DejaVu from `fonts-dejavu-core` (see backend/Dockerfile); the rest are common
# developer machines, so `pytest` exercises the real PDF path outside Docker too.
FONT_CANDIDATES: tuple[Path, ...] = (
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    Path("/usr/local/share/fonts/DejaVuSans.ttf"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
)
# Bold companions are looked up next to the regular face by filename.
BOLD_SUFFIXES: tuple[tuple[str, str], ...] = (
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
    ("Arial.ttf", "Arial Bold.ttf"),
)

PDF_FONT_NAME = "TranscriptSans"
PDF_BOLD_FONT_NAME = "TranscriptSans-Bold"


def format_timestamp(seconds: float | None) -> str:
    """Seconds on the global timeline -> `00:12:34`."""
    total = max(0, int(seconds or 0))
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def speaker_names(transcript: Transcript) -> dict[str, str]:
    return {speaker.id: speaker.display_name for speaker in transcript.speakers}


def segment_heading(transcript: Transcript, segment) -> str:
    """`[00:00:12 – 00:00:18] Спикер 1` (SPEC.md §6)."""
    names = speaker_names(transcript)
    name = names.get(segment.speaker_id or "", UNKNOWN_SPEAKER)
    start = format_timestamp(segment.start)
    end = format_timestamp(segment.end)
    return f"[{start}{TIME_RANGE_SEPARATOR}{end}] {name}"


# --- T6.1 TXT ------------------------------------------------------------------------


def render_txt(transcript: Transcript) -> str:
    blocks = [
        f"{segment_heading(transcript, segment)}\n{segment.text}" for segment in transcript.segments
    ]
    return "\n\n".join(blocks) + "\n" if blocks else ""


def write_txt(transcript: Transcript, destination: Path) -> Path:
    destination.write_text(render_txt(transcript), encoding="utf-8")
    return destination


# --- T6.2 DOCX -----------------------------------------------------------------------


def write_docx(transcript: Transcript, destination: Path) -> Path:
    document = Document()
    document.add_heading(DOCUMENT_TITLE, level=0)

    document.add_paragraph(f"Файл: {transcript.source.filename}")
    if transcript.source.duration_seconds is not None:
        document.add_paragraph(
            f"Длительность: {format_timestamp(transcript.source.duration_seconds)}"
        )

    for segment in transcript.segments:
        heading = document.add_paragraph()
        run = heading.add_run(segment_heading(transcript, segment))
        run.bold = True
        run.font.size = Pt(10)
        document.add_paragraph(segment.text)

    document.save(str(destination))
    return destination


# --- T6.3 PDF ------------------------------------------------------------------------


def _bold_companion(regular: Path) -> Path | None:
    for regular_name, bold_name in BOLD_SUFFIXES:
        if regular.name == regular_name:
            candidate = regular.with_name(bold_name)
            if candidate.is_file():
                return candidate
    return None


def resolve_pdf_font(configured: Path | None = None) -> Path:
    """Find a Cyrillic-capable TTF. Built-in PDF fonts cover Latin only (TASKS.md T6.3)."""
    if configured is not None:
        if not Path(configured).is_file():
            raise AppError(
                ErrorCode.EXPORT_FAILED,
                f"PDF_FONT_PATH points at {configured}, which does not exist",
            )
        return Path(configured)

    for candidate in FONT_CANDIDATES:
        if candidate.is_file():
            return candidate

    raise AppError(
        ErrorCode.EXPORT_FAILED,
        "No Cyrillic-capable font was found for the PDF export. Install DejaVu Sans "
        "or point PDF_FONT_PATH at a TTF file.",
    )


def _register_pdf_font(font_path: Path) -> str:
    """Register the font once per process and return the bold face name to use."""
    registered = pdfmetrics.getRegisteredFontNames()
    if PDF_FONT_NAME not in registered:
        pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(font_path)))

    bold_path = _bold_companion(font_path)
    if bold_path is None:
        # Better a non-bold heading than a failed export.
        return PDF_FONT_NAME
    if PDF_BOLD_FONT_NAME not in registered:
        pdfmetrics.registerFont(TTFont(PDF_BOLD_FONT_NAME, str(bold_path)))
    return PDF_BOLD_FONT_NAME


def write_pdf(transcript: Transcript, destination: Path, font_path: Path | None = None) -> Path:
    resolved = resolve_pdf_font(font_path)
    bold_name = _register_pdf_font(resolved)

    body = ParagraphStyle(
        "TranscriptBody",
        fontName=PDF_FONT_NAME,
        fontSize=10,
        leading=14,
        alignment=TA_LEFT,
    )
    heading = ParagraphStyle(
        "TranscriptHeading", parent=body, fontName=bold_name, spaceBefore=8, spaceAfter=2
    )
    title = ParagraphStyle("TranscriptTitle", parent=body, fontName=bold_name, fontSize=16)

    flow = [Paragraph(escape(DOCUMENT_TITLE), title), Spacer(1, 6)]
    flow.append(Paragraph(escape(f"Файл: {transcript.source.filename}"), body))
    if transcript.source.duration_seconds is not None:
        duration = format_timestamp(transcript.source.duration_seconds)
        flow.append(Paragraph(escape(f"Длительность: {duration}"), body))

    for segment in transcript.segments:
        flow.append(Paragraph(escape(segment_heading(transcript, segment)), heading))
        flow.append(Paragraph(escape(segment.text), body))

    SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        title=DOCUMENT_TITLE,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    ).build(flow)
    return destination


# --- T6.4 regeneration ---------------------------------------------------------------


def export_path(result_dir: Path, export_format: str) -> Path:
    return result_dir / f"transcript.{export_format}"


def load_transcript(path: Path) -> Transcript:
    """The canonical JSON is the only input an export is ever built from."""
    if not path.is_file():
        raise AppError(ErrorCode.EXPORT_FAILED, "Canonical transcript.json is missing")
    try:
        return Transcript.model_validate_json(path.read_text(encoding="utf-8"))
    except ValueError as exc:
        logger.exception("could not read transcript.json", extra={"stage": "exports"})
        raise AppError(ErrorCode.EXPORT_FAILED, "Canonical transcript.json is unreadable") from exc


def generate_exports(
    transcript: Transcript,
    result_dir: Path,
    font_path: Path | None = None,
) -> list[Path]:
    """Write every view next to the canonical JSON. Existing files are overwritten."""
    result_dir.mkdir(parents=True, exist_ok=True)
    writers = {
        "txt": lambda path: write_txt(transcript, path),
        "docx": lambda path: write_docx(transcript, path),
        "pdf": lambda path: write_pdf(transcript, path, font_path),
    }

    written: list[Path] = []
    for export_format in EXPORT_FORMATS:
        destination = export_path(result_dir, export_format)
        try:
            written.append(writers[export_format](destination))
        except AppError:
            raise
        except Exception as exc:
            logger.exception("export failed", extra={"stage": "exports", "format": export_format})
            destination.unlink(missing_ok=True)
            raise AppError(
                ErrorCode.EXPORT_FAILED, f"Could not generate the {export_format.upper()} export"
            ) from exc

    logger.info(
        "exports generated",
        extra={
            "stage": "exports",
            "formats": list(EXPORT_FORMATS),
            "segments": len(transcript.segments),
        },
    )
    return written


def regenerate_exports(
    transcript_path: Path,
    result_dir: Path,
    font_path: Path | None = None,
) -> list[Path]:
    """T6.4 — rebuild TXT/DOCX/PDF from the canonical JSON on disk."""
    return generate_exports(load_transcript(transcript_path), result_dir, font_path)
